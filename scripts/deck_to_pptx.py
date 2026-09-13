#!/usr/bin/env python3
"""Turn the HTML deck into a PPTX that opens anywhere, without losing the design.

The deck is a web page: IBM Plex, hairline rules, a day bar drawn from real duty segments.
Converting that into PowerPoint shapes would mean rebuilding it badly. So each slide goes in as
a full-bleed image of our own render — pixel-identical — and the presenter notes ride along in
the notes field, which is the one thing the PDF cannot carry.

Not editable as text. If a line needs changing, change docs/demo/deck/index.html and run this
again; the HTML stays the source.

    python3 scripts/deck_to_pptx.py            # -> docs/demo/deck/dockrisk-deck.pptx
"""
from __future__ import annotations

import html
import json
import pathlib
import re
import subprocess
import sys
import tempfile

ROOT = pathlib.Path(__file__).resolve().parent.parent
DECK = ROOT / "docs" / "demo" / "deck" / "index.html"
OUT = ROOT / "docs" / "demo" / "deck" / "dockrisk-deck.pptx"
PDF = ROOT / "docs" / "demo" / "deck" / "dockrisk-deck.pdf"
SHOT = ROOT / "apps" / "web" / "scripts" / "_pptx_shot.mjs"
W, H = 1920, 1200          # 16:10, the deck's own ratio at 1440x900

RENDERER = r"""
import puppeteer from "puppeteer-core";
import { mkdirSync, readFileSync, writeFileSync } from "node:fs";
const [url, outdir, w, h, posterPath] = process.argv.slice(2);
mkdirSync(outdir, { recursive: true });
const b = await puppeteer.launch({
  executablePath: "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
  headless: true, args: ["--no-sandbox", "--disable-gpu", "--hide-scrollbars", "--font-render-hinting=none"],
});
const rects = {};
const poster = posterPath && posterPath !== "-"
  ? "data:image/png;base64," + readFileSync(posterPath).toString("base64") : null;
const p = await b.newPage();
await p.setViewport({ width: Number(w), height: Number(h), deviceScaleFactor: 1 });
await p.goto(url, { waitUntil: "networkidle2", timeout: 60000 });
await p.evaluate(() => document.fonts.ready);
// Kill every transition and force the revealed state. Adding .visible and waiting caught
// staggered fades mid-way and baked half-transparent figures into the slides — the deck's own
// print stylesheet does exactly this for the same reason.
await p.addStyleTag({ content: `
  *, *::before, *::after { transition: none !important; animation: none !important; }
  .reveal, .reveal.visible { opacity: 1 !important; transform: none !important; }
  .daybar .seg { width: var(--w) !important; }
  .daybar .tick, .figures .figure { opacity: 1 !important; transform: none !important; }
` });
await new Promise(r => setTimeout(r, 2000));
const n = await p.evaluate(() => document.querySelectorAll(".slide").length);
for (let i = 0; i < n; i++) {
  // the deck scrolls its body, not the window
  await p.evaluate((i) => {
    const s = document.querySelectorAll(".slide")[i];
    document.body.style.scrollBehavior = "auto";
    document.body.scrollTop = s.offsetTop;
    s.classList.add("visible");
    s.querySelectorAll(".reveal").forEach(r => r.classList.add("visible"));
    for (const el of document.querySelectorAll(".nav-dots,.keyboard-hint,.progress-bar,.edit-toggle,.edit-bar,.notes-panel")) el.style.display = "none";
  }, i);
  await new Promise(r => setTimeout(r, 500));
  // A <video> screenshots with its browser chrome showing — a painted play button, a painted
  // 0:00 / 1:30, a painted progress bar. Baked into a slide that is itself a picture, that reads
  // as a video that will not play, and the first thing anyone does is click the dead pixel. So
  // record where the video sat, then swap it for a real frame of the film before the shot; the
  // PPTX puts the movie back in exactly that rectangle with the same frame as its poster.
  const v = await p.evaluate((i, poster) => {
    const s = document.querySelectorAll(".slide")[i];
    const el = s.querySelector("video");
    if (!el) return null;
    const r = el.getBoundingClientRect(), sr = s.getBoundingClientRect();
    const rect = { src: el.getAttribute("src"),
                   x: (r.left - sr.left) / sr.width, y: (r.top - sr.top) / sr.height,
                   w: r.width / sr.width, h: r.height / sr.height };
    if (poster) {
      // A <video> with a poster and no controls paints the frame and nothing else — no play
      // button, no scrubber, no duration. Keeping the element keeps the slide's own CSS, which
      // swapping in an <img> did not.
      el.removeAttribute("controls");
      el.setAttribute("poster", poster);
      el.load();
    }
    return rect;
  }, i, poster);
  if (v) {
    rects[i] = v;
    await new Promise(r => setTimeout(r, 400));   // let the swapped-in frame decode
  }
  await p.screenshot({ path: `${outdir}/s${String(i).padStart(2, "0")}.png` });
}
writeFileSync(`${outdir}/videos.json`, JSON.stringify(rects));
console.log(n);
await b.close();
"""


POSTER_AT = "3"          # far enough in that the board is populated, before the first cut


def poster_frame(tmp: pathlib.Path) -> pathlib.Path | None:
    """One frame of the fallback film, used twice: painted into the slide picture where the
    <video> was, and handed to PowerPoint as the movie's poster. Same frame both times, so the
    slide looks identical whether the movie loads or not — and nothing on it looks clickable
    that is not."""
    src = DECK.parent / "assets" / "backup.mp4"
    if not src.exists():
        return None
    out = tmp / "poster.png"
    r = subprocess.run(["ffmpeg", "-nostdin", "-loglevel", "error", "-y", "-ss", POSTER_AT,
                        "-i", str(src), "-frames:v", "1", str(out)], capture_output=True)
    return out if r.returncode == 0 and out.exists() else None


def notes_per_slide() -> list[str]:
    """Presenter notes, in slide order, stripped to plain text."""
    src = DECK.read_text(encoding="utf-8")
    out = []
    for chunk in re.split(r'(?=<section class="slide)', src):
        if 'class="slide' not in chunk:
            continue
        found = re.findall(r'class="notes[^"]*"[^>]*>(.*?)</(?:div|aside|section)>', chunk, re.S)
        text = " ".join(found)
        # PowerPoint's notes pane is plain text, so the two kinds of line have to stay apart by
        # shape instead of by style: the sentences he speaks get an arrow and their own line,
        # the directions he only reads sit flush left.
        text = re.sub(r'<p class="say"[^>]*>', "\n▶ ", text)
        text = re.sub(r'<p class="saylabel"[^>]*>', "\n\n", text)
        text = re.sub(r'<p class="dir"[^>]*>', "\n\n", text)
        text = re.sub(r"<li[^>]*>", "\n• ", text)
        text = re.sub(r"<br\s*/?>", "\n", text)
        text = html.unescape(re.sub(r"<[^>]+>", " ", text))
        text = re.sub(r"[ \t]+", " ", text)
        out.append(re.sub(r"\n{3,}", "\n\n", "\n".join(l.strip() for l in text.split("\n"))).strip())
    return out


def write_pdf(pngs: list[pathlib.Path]) -> None:
    try:
        from PIL import Image
    except ImportError:
        print("  (no Pillow — skipping the PDF)", file=sys.stderr)
        return
    pages = [Image.open(p).convert("RGB") for p in pngs]
    pages[0].save(PDF, save_all=True, append_images=pages[1:], resolution=150.0)
    print(f"{len(pages)} slides · {PDF.stat().st_size / 1e6:.1f} MB -> {PDF.relative_to(ROOT)}")


def main() -> int:
    from pptx import Presentation
    from pptx.util import Emu

    SHOT.write_text(RENDERER, encoding="utf-8")
    tmp = pathlib.Path(tempfile.mkdtemp(prefix="deckpng-"))
    try:
        poster = poster_frame(tmp)
        r = subprocess.run(["node", str(SHOT), DECK.as_uri(), str(tmp), str(W), str(H),
                            str(poster) if poster else "-"],
                           cwd=SHOT.parent.parent, capture_output=True, text=True)
        if r.returncode != 0:
            print(r.stderr[-800:], file=sys.stderr)
            return 1
        pngs = sorted(tmp.glob("s*.png"))
        notes = notes_per_slide()
        if len(pngs) != len(notes):
            print(f"!! {len(pngs)} rendered vs {len(notes)} sets of notes — check the deck markup", file=sys.stderr)

        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(7620000)   # 16:10 at PowerPoint scale
        blank = prs.slide_layouts[6]
        vids = {}
        vf = tmp / "videos.json"
        if vf.exists():
            vids = {int(k): v for k, v in json.loads(vf.read_text()).items()}

        for i, png in enumerate(pngs):
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
            v = vids.get(i)
            if v:
                mp4 = (DECK.parent / v["src"]).resolve()
                if mp4.exists():
                    s.shapes.add_movie(str(mp4),
                                       int(v["x"] * prs.slide_width), int(v["y"] * prs.slide_height),
                                       int(v["w"] * prs.slide_width), int(v["h"] * prs.slide_height),
                                       poster_frame_image=str(poster) if poster else None,
                                       mime_type="video/mp4")
                    print(f"  slide {i+1}: embedded {mp4.name}"
                          f"{' with a real poster frame' if poster else ' (no poster — grey box in PowerPoint)'}")
                else:
                    print(f"  !! slide {i+1}: {mp4} not found", file=sys.stderr)
            note = notes[i] if i < len(notes) else ""
            if note:
                s.notes_slide.notes_text_frame.text = note
        prs.save(OUT)
        size = OUT.stat().st_size / 1e6
        with_notes = sum(1 for n in notes if n)
        print(f"{len(pngs)} slides · notes on {with_notes} · {size:.1f} MB -> {OUT.relative_to(ROOT)}")

        # The PDF comes off the same renders, so the two can never disagree about what a slide
        # looks like. It carries no notes — that is what the PPTX is for.
        write_pdf(pngs)
        return 0
    finally:
        SHOT.unlink(missing_ok=True)
        for f in tmp.glob("*"):
            f.unlink()
        tmp.rmdir()


if __name__ == "__main__":
    raise SystemExit(main())
