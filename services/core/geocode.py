"""Geocode the cities and postal codes that appear in the organizer workbook.

Nominatim, 1 request/second, results cached to data/geo/. Run once:
    uv run python -m core.geocode
"""
from __future__ import annotations

import json
import re
import sys
import time
from pathlib import Path

import httpx
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "data/portal-downloads/Hackathon_Data.xlsx"
GEO = ROOT / "data/geo"
UA = "roadstar-hackathon-dispatch/0.1 (hackathon project; contact via portal)"
NOMINATIM = "https://nominatim.openstreetmap.org/search"

PROV_OK = {"ON"}


def norm_city(desc: str | None) -> tuple[str, str] | None:
    """'MILTON,ON' / 'MILTON, ON' / 'ST CATHARINES, ON' -> ('MILTON', 'ON')."""
    if not desc:
        return None
    m = re.match(r"\s*(.+?)\s*,\s*([A-Z]{2})\s*$", str(desc).strip().upper())
    if not m:
        return None
    return m.group(1).strip(), m.group(2)


def collect_targets() -> tuple[set[tuple[str, str]], set[str]]:
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    cities: set[tuple[str, str]] = set()
    postals: set[str] = set()

    ws = wb["Dispatch"]
    rows = ws.iter_rows(values_only=True)
    h = [str(x) for x in next(rows)]
    ix = {c: i for i, c in enumerate(h)}
    zone_cols = [ix[c] for c in ("LEGO_ZONE_DESC", "LEGD_ZONE_DESC", "ORIG_ZONE_DESC", "DEST_ZONE_DESC")]
    for r in rows:
        for c in zone_cols:
            t = norm_city(r[c])
            if t and t[1] in PROV_OK:
                cities.add(t)

    ws = wb["Tlorder"]
    rows = ws.iter_rows(values_only=True)
    h = [str(x) for x in next(rows)]
    ix = {c: i for i, c in enumerate(h)}
    for r in rows:
        for pc_col, prov_col, city_col in (("ORIGPC", "ORIGPROV", "ORIGCITY"), ("DESTPC", "DESTPROV", "DESTCITY")):
            if r[ix[prov_col]] in PROV_OK:
                pc = str(r[ix[pc_col]] or "").replace(" ", "").upper()
                if re.fullmatch(r"[A-Z]\d[A-Z]\d[A-Z]\d", pc):
                    postals.add(pc)
                city = str(r[ix[city_col]] or "").strip().upper()
                if city:
                    cities.add((city, r[ix[prov_col]]))
    return cities, postals


def lookup(client: httpx.Client, params: dict) -> dict | None:
    for attempt in range(3):
        try:
            resp = client.get(NOMINATIM, params={**params, "format": "json", "limit": 1, "countrycodes": "ca"})
            if resp.status_code == 200:
                data = resp.json()
                if data:
                    d = data[0]
                    return {"lat": float(d["lat"]), "lon": float(d["lon"]), "display": d.get("display_name", "")}
                return None
        except httpx.HTTPError:
            pass
        time.sleep(2 * (attempt + 1))
    return None


def main() -> None:
    GEO.mkdir(parents=True, exist_ok=True)
    cities_path, postal_path = GEO / "cities.json", GEO / "postal.json"
    cities_cache = json.loads(cities_path.read_text()) if cities_path.exists() else {}
    postal_cache = json.loads(postal_path.read_text()) if postal_path.exists() else {}

    cities, postals = collect_targets()
    print(f"targets: {len(cities)} cities, {len(postals)} postal codes; cached {len(cities_cache)}/{len(postal_cache)}", flush=True)

    with httpx.Client(headers={"User-Agent": UA}, timeout=20) as client:
        for city, prov in sorted(cities):
            key = f"{city}, {prov}"
            if key in cities_cache:
                continue
            cities_cache[key] = lookup(client, {"city": city.title(), "state": "Ontario"})
            cities_path.write_text(json.dumps(cities_cache, indent=1))
            print(f"city {key}: {'ok' if cities_cache[key] else 'MISS'}", flush=True)
            time.sleep(1.1)
        for pc in sorted(postals):
            if pc in postal_cache:
                continue
            postal_cache[pc] = lookup(client, {"postalcode": f"{pc[:3]} {pc[3:]}"})
            postal_path.write_text(json.dumps(postal_cache, indent=1))
            print(f"postal {pc}: {'ok' if postal_cache[pc] else 'MISS'}", flush=True)
            time.sleep(1.1)

    print("done", flush=True)


if __name__ == "__main__":
    sys.exit(main())
